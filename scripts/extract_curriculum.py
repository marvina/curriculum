import json
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "0001视觉学院数据_已修正.xlsx"
OUTPUT = Path(__file__).resolve().parents[1] / "lib" / "curriculum-data.json"
MAJOR_NAMES = {
    "视觉传达": "视觉传达设计",
    "数字媒体": "数字媒体艺术",
    "包装设计": "包装设计",
    "智能交互": "智能交互设计",
}


workbook = openpyxl.load_workbook(SOURCE, data_only=True, read_only=True)
courses = []

for sheet in workbook.worksheets:
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    for values in sheet.iter_rows(min_row=2, values_only=True):
        row = dict(zip(headers, values))
        if not row.get("课程名称"):
            continue
        total_hours = row.get("总学时") or 0
        practice_hours = row.get("实践学时") or 0
        courses.append({
            "id": f"{sheet.title}-{len(courses) + 1}",
            "major": MAJOR_NAMES[sheet.title],
            "module": row.get("课程模组"),
            "system": row.get("课程体系"),
            "group": row.get("课程种类"),
            "nature": row.get("课程性质"),
            "name": row.get("课程名称"),
            "credits": row.get("学分"),
            "semester": row.get("上课学期"),
            "hours": {
                "total": total_hours,
                "theory": row.get("理论学时") or 0,
                "practice": practice_hours,
            },
            "practiceRate": round(practice_hours / total_hours * 100) if total_hours else 0,
            "weeklyHours": row.get("周课时"),
            "weeks": row.get("上课周数"),
            "note": row.get("备注"),
        })

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
OUTPUT.write_text(json.dumps({"version": "2026", "courses": courses}, ensure_ascii=False, indent=2), encoding="utf-8")
print(f"Wrote {len(courses)} courses to {OUTPUT}")
